import docx

doc = docx.Document(r'D:\NIC\eWork3.0\Role_Master_API_Frontend_Integration_Guide.docx')
print("=== PARAGRAPHS ===")
for p in doc.paragraphs:
    if p.text.strip():
        print(p.text)

print("\n=== TABLES ===")
for i, table in enumerate(doc.tables):
    print(f"--- Table {i+1} ---")
    for row in table.rows:
        print([cell.text.strip().replace('\n', ' ') for cell in row.cells])
